"""Lịch sử nạp, Chi tiết lần nạp, Đối chiếu, Tệp gốc, Gỡ / Xoá lịch sử (mục 16–18)."""

from __future__ import annotations

import csv
import io

import openpyxl
import pytest

from .conftest import can_tep_mau, dem, tai_len


def _nap(goi, nam="2026", tep=None):
    ma = tai_len(goi, tep or can_tep_mau(), nam=nam).json()["pending_id"]
    return goi("POST", f"/uploads/{ma}/confirm").json()["load_id"]


@pytest.fixture
def mot_lan(admin):
    return _nap(admin)


def test_danh_sach_va_loc(admin, loader, mot_lan, tmp_path):
    so = openpyxl.load_workbook(can_tep_mau())
    del so["TỔNG HỢP"]
    tep = tmp_path / "loi.xlsx"
    so.save(tep)
    tu_choi = tai_len(loader, tep, nam="2027").json()["load_id"]

    d = admin("GET", "/loads?domain=HQKD").json()
    assert d["total"] == 2
    assert d["uploaders"] == ["admin", "loader"]
    assert [x["id"] for x in d["items"]] == [tu_choi, mot_lan]
    dong = d["items"][1]
    assert dong["file_name"] == can_tep_mau().name
    assert (dong["year"], dong["user"], dong["row_count"], dong["status"]) == (
        2026, "admin", 7561, "success")
    assert dong["file_type"] == {"code": "BAO_CAO_HQKH", "name": "Báo cáo hiệu quả từng khách hàng"}
    assert d["items"][0]["row_count"] == 0 and d["items"][0]["status"] == "rejected"

    def ids(q):
        return [x["id"] for x in admin("GET", "/loads?domain=HQKD&" + q).json()["items"]]

    assert ids("year=2026") == [mot_lan]
    assert ids("status=rejected") == [tu_choi]
    assert ids("user=loader") == [tu_choi]
    assert ids("query=loi.xlsx") == [tu_choi]
    assert ids(f"query=%23{mot_lan}") == [mot_lan]
    assert ids("file_type=BAO_CAO_HQKH&page_size=25&page=1") == [tu_choi, mot_lan]
    assert ids("page_size=25&page=2") == []
    assert admin("GET", "/loads?domain=HQKD&page_size=10").status_code == 422
    assert admin("GET", "/loads?domain=HQKD&status=xyz").status_code == 422
    assert admin("GET", "/loads?domain=HQ-MAU-GC").json()["total"] == 0


def test_chi_tiet_thanh_cong(admin, mot_lan):
    d = admin("GET", f"/loads/{mot_lan}").json()
    assert (d["id"], d["status"], d["domain"], d["year"]) == (mot_lan, "success", "HQKD", 2026)
    assert d["file_type"]["code"] == "BAO_CAO_HQKH"
    assert d["months"] == "1 → 7"
    assert (d["total_rows"], d["rows_written"], d["sheet_count"]) == (7561, 7471, 3)
    assert d["sheet"] is None and d["errors"] == []
    assert d["can_rollback"] and d["is_active"]
    assert d["primary_table"] == "fact_ket_qua_kd"
    assert [b["code"] for b in d["steps"]] == ["A1", "A2", "A3", "A4", "A5",
                                               "B1", "B2", "B3", "B4", "B5"]
    assert all(b["status"] == "ok" for b in d["steps"])
    buoc = {b["code"]: b for b in d["steps"]}
    assert buoc["B1"]["result"] == ("Đọc lại tệp độc lập rồi so: 7.561 dòng trong tệp = "
                                     "7.561 dòng đã ghi.")
    assert buoc["B2"]["result"] == "So với số dòng sẽ ghi ở A3: 7.471 = 7.471 dòng."
    assert buoc["B3"]["result"] == "7.471 = 7.471 dòng."
    assert buoc["B4"]["result"] == (
        "Lệch 0 ở mọi tháng, mọi cột tiền; số tháng, khách hàng, nhóm, khoản mục khớp. "
        "Tổng cả bảng lệch 0 ở 2/2 bảng (Chi phí 122.689.472.100,65; Kết quả kinh doanh "
        "3.193.114.446.744,50).")
    assert buoc["B5"]["result"] == "Dữ liệu tháng 1–7 năm 2026 có hiệu lực từ lúc này."
    assert [b["verdict"] for b in d["steps"][5:]] == ["Khớp", "Khớp", "Khớp", "Đúng",
                                                      "Đã chốt"]
    assert admin("GET", "/loads/999999").status_code == 404


def test_doi_chieu_hqkd_5_the(admin, mot_lan):
    the = admin("GET", f"/loads/{mot_lan}/reconcile").json()["cards"]
    assert [t["key"] for t in the] == ["row_count", "totals", "amount_by_group", "codes", "before_after"]
    assert [t["title"] for t in the] == [
        "Đối chiếu số dòng: tệp gốc ↔ database",
        "Đối chiếu tổng tiền cả bảng: tệp gốc ↔ database",
        "Đối chiếu tổng tiền theo tháng, cột tiền, nhóm",
        "Đối chiếu tháng, khách hàng, nhóm, khoản mục",
        "Các tháng trong năm 2026 trước và sau lần nạp",
    ]
    so_dong = the[0]
    assert so_dong["totals"] == ["Tổng", 7561, 7561, 7471, 7471, 90, None]
    kh = so_dong["rows"][2]
    assert kh["cells"][0] == "Danh mục khách hàng" and kh["note"] == "dim_khach_hang"
    assert kh["verdict_note"] == "Đủ sau khi bỏ 90 dòng trùng y hệt"
    assert kh["cells"][1] == {"value": 227, "link": {"sheet": 2}}
    assert kh["cells"][3]["link"] == {"table": "dim_khach_hang", "layer": "silver"}
    assert [d["cells"][1:] for d in the[1]["rows"]] == [
        [122689472100.6453, 122689472100.6453, 0],
        [3193114446744.495, 3193114446744.495, 0]]
    tien = the[2]
    assert tien["options"][0]["value"] == "fact_ket_qua_kd"
    assert tien["options"][1]["value"] == "doanh_thu"
    assert tien["rows"][0]["cells"][:4] == ["Tháng 1", 50, 50, 196631400595]
    assert tien["rows"][2]["link"] == {"table": "fact_ket_qua_kd", "layer": "gold", "year": 2026,
                                       "period": 3}
    ma = {d["cells"][0]: d for d in the[3]["rows"]}
    assert ma["Số khách hàng trong Kết quả kinh doanh"]["cells"][1:] == [69, 69]
    assert ma["Số khoản mục trong Chi phí"]["cells"][1:] == [20, 20]
    assert ma["Số mã khách hàng trong Danh mục khách hàng"]["cells"][1:] == [
        "135 mã (227 dòng)", "135 mã (137 dòng)"]
    thieu = ma["Mã khách hàng có trong số liệu nhưng không có trong Danh mục khách hàng"]
    assert thieu["cells"][1:] == ["BP-TTXK-N04 · TARGET AUSTRALIA PTY. LTD. (tháng 6, 7)", 1]
    assert thieu["verdict"] == "Ghi nhận"
    truoc_sau = the[4]
    assert len(truoc_sau["rows"]) == 12
    assert truoc_sau["rows"][0]["cells"] == ["Tháng 1/2026", "Có", "0 / 0", "50 / 1.000",
                                         "Ghi thêm"]
    assert truoc_sau["rows"][11]["verdict"] == "Giữ nguyên"


def test_doi_chieu_mot_the_theo_lua_chon(admin, mot_lan):
    url = (f"/loads/{mot_lan}/reconcile?card=amount_by_group&table=fact_chi_phi"
           "&column=so_tien&group_by=ma_khoan_cp")
    the = admin("GET", url).json()["cards"]
    assert len(the) == 1
    t = the[0]
    assert t["columns"][0]["label"] == "Khoản mục"
    assert [c["value"] for c in t["options"][2]["choices"]] == [
        "ky_thang", "ma_nhom_kd", "ma_khoan_cp"]
    assert any(d.get("note") for d in t["rows"])
    assert t["totals"][1] == 6960 and t["totals"][3] == t["totals"][4]
    # Tổng các nhóm bằng tổng cả bảng (NT-15).
    tong_bang = admin("GET", f"/loads/{mot_lan}/reconcile?card=totals").json()["cards"][0]
    assert t["totals"][3] == tong_bang["rows"][0]["cells"][1]
    kq = admin("GET", f"/loads/{mot_lan}/reconcile?card=amount_by_group&table=fact_ket_qua_kd"
                      "&column=all").json()["cards"][0]
    tong_kqkd = tong_bang["rows"][1]["cells"][1]
    assert kq["totals"][3] == pytest.approx(tong_kqkd)
    assert admin("GET", f"/loads/{mot_lan}/reconcile?card=khong_co").status_code == 404


def test_tai_danh_sach_loi_csv(admin, tmp_path):
    so = openpyxl.load_workbook(can_tep_mau())
    del so["DANH MỤC CHI PHÍ"]
    tep = tmp_path / "loi.xlsx"
    so.save(tep)
    lid = tai_len(admin, tep).json()["load_id"]
    r = admin("GET", f"/loads/{lid}/errors.csv")
    assert r.status_code == 200
    assert r.headers["content-disposition"] == f'attachment; filename="loi-lan-nap-{lid}.csv"'
    dong = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert dong[0] == ["Sheet", "Vị trí", "Vấn đề", "Cách xử lý"]
    assert dong[1][:2] == ["DANH MỤC CHI PHÍ", "Toàn bộ sheet"]
    the = admin("GET", f"/loads/{lid}/reconcile").json()["cards"]
    assert the == [{"key": "rejected", "title": "Đối chiếu tệp gốc ↔ database",
                    "message": "Không có dữ liệu đối chiếu vì tệp bị từ chối ở bước "
                               "kiểm tra cấu trúc, chưa có gì được ghi vào database.",
                    "columns": [], "rows": []}]


def test_tep_goc(admin, mot_lan):
    r = admin("GET", f"/loads/{mot_lan}/file")
    assert r.status_code == 200 and r.content == can_tep_mau().read_bytes()
    sheet = admin("GET", f"/loads/{mot_lan}/file/sheets").json()
    assert [(s["index"], s["name"], s["row_count"]) for s in sheet] == [
        (1, "TỔNG HỢP", 524), (2, "DANH SACH KHACH HANG THEO HDX", 232),
        (3, "DANH MỤC CHI PHÍ", 27)]
    nd = admin("GET", f"/loads/{mot_lan}/file/sheets/3?page=1&page_size=25").json()
    assert nd["total"] == 27 and len(nd["rows"]) == 25 and nd["rows"][0]["row_number"] == 1
    assert set(nd["rows"][0]) == {"row_number", "cells", "hidden"}
    assert admin("GET", f"/loads/{mot_lan}/file/sheets/9").status_code == 404


def test_go_lan_khong_moi_nhat_409(admin, loader, sach):
    """QT-17, B18: chỉ gỡ được lần nạp thành công mới nhất; gỡ xong dữ liệu lần
    trước được dùng lại."""
    lan1 = _nap(admin)
    lan2 = _nap(admin)
    assert dem(sach, "SELECT count(*) FROM gold.fact_ket_qua_kd WHERE is_current") == 348
    assert loader("DELETE", f"/loads/{lan2}").status_code == 403
    r = admin("DELETE", f"/loads/{lan1}")
    assert r.status_code == 409
    assert r.json()["detail"] == ("Có 1 lần nạp sau lần này trên cùng nhóm thông tin. "
                                  "Hãy gỡ các lần nạp sau trước.")
    assert admin("GET", f"/loads/{lan1}").json()["can_rollback"] is False
    assert admin("DELETE", f"/loads/{lan2}").status_code == 204
    assert admin("GET", f"/loads/{lan2}").status_code == 404
    assert dem(sach, "SELECT count(*) FROM gold.fact_ket_qua_kd WHERE is_current") == 348
    assert admin("DELETE", f"/loads/{lan1}").status_code == 204
    assert dem(sach, "SELECT count(*) FROM gold.fact_ket_qua_kd") == 0
    with sach.catalog.transaction() as so:
        hanh_dong = [r[0] for r in so.execute(
            "SELECT action FROM ctl_audit_event WHERE object_id IN (?, ?) ORDER BY event_id",
            (str(lan1), str(lan2)))]
    assert hanh_dong.count("load.success") == 2
    assert hanh_dong.count("load.rollback") == 2


def test_xoa_lich_su_lan_bi_tu_choi(admin, sach, tmp_path):
    so = openpyxl.load_workbook(can_tep_mau())
    del so["DANH MỤC CHI PHÍ"]
    tep = tmp_path / "loi.xlsx"
    so.save(tep)
    lid = tai_len(admin, tep).json()["load_id"]
    assert admin("DELETE", f"/loads/{lid}").status_code == 204
    assert dem(sach, "SELECT count(*) FROM ctl.load") == 0
    assert admin("DELETE", f"/loads/{lid}").status_code == 404
