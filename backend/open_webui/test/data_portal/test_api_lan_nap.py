"""Lịch sử nạp, Chi tiết lần nạp, Đối chiếu, Tệp gốc, Gỡ / Xoá lịch sử (mục 16–18)."""

from __future__ import annotations

import csv
import io

import openpyxl
import pytest

from .conftest import can_tep_mau, dem, tai_len


def _nap(goi, nam="2026", tep=None):
    ma = tai_len(goi, tep or can_tep_mau(), nam=nam).json()["ma_tep_cho"]
    return goi("POST", f"/uploads/{ma}/confirm").json()["load_id"]


@pytest.fixture
def mot_lan(admin):
    return _nap(admin)


def test_danh_sach_va_loc(admin, loader, mot_lan, tmp_path):
    so = openpyxl.load_workbook(can_tep_mau())
    del so["TỔNG HỢP"]
    tep = tmp_path / "loi.xlsx"
    so.save(tep)
    tu_choi = tai_len(loader, tep, nam="2027").json()["load_id"]

    d = admin("GET", "/loads?nhom=HQKD").json()
    assert d["tong"] == 2
    assert d["nguoi_nap"] == ["admin", "loader"]
    assert [x["id"] for x in d["dong"]] == [tu_choi, mot_lan]
    dong = d["dong"][1]
    assert dong["ten_tep"] == can_tep_mau().name
    assert (dong["nam"], dong["nguoi"], dong["so_dong"], dong["status"]) == (
        2026, "admin", 7561, "success")
    assert dong["loai"] == {"ma": "BAO_CAO_HQKH", "ten": "Báo cáo hiệu quả từng khách hàng"}
    assert d["dong"][0]["so_dong"] == 0 and d["dong"][0]["status"] == "rejected"

    def ids(q):
        return [x["id"] for x in admin("GET", "/loads?nhom=HQKD&" + q).json()["dong"]]

    assert ids("nam=2026") == [mot_lan]
    assert ids("trang_thai=rejected") == [tu_choi]
    assert ids("nguoi=loader") == [tu_choi]
    assert ids("tim=loi.xlsx") == [tu_choi]
    assert ids(f"tim=%23{mot_lan}") == [mot_lan]
    assert ids("loai=BAO_CAO_HQKH&moi=25&trang=1") == [tu_choi, mot_lan]
    assert ids("moi=25&trang=2") == []
    assert admin("GET", "/loads?nhom=HQKD&moi=10").status_code == 422
    assert admin("GET", "/loads?nhom=HQKD&trang_thai=xyz").status_code == 422
    assert admin("GET", "/loads?nhom=HQ-MAU-GC").json()["tong"] == 0


def test_chi_tiet_thanh_cong(admin, mot_lan):
    d = admin("GET", f"/loads/{mot_lan}").json()
    assert (d["id"], d["status"], d["nhom"], d["nam"]) == (mot_lan, "success", "HQKD", 2026)
    assert d["loai"]["ma"] == "BAO_CAO_HQKH"
    assert d["thang"] == "1 → 7"
    assert (d["tong_so_dong"], d["so_dong_ghi"], d["so_sheet"]) == (7561, 7471, 3)
    assert d["sheet"] is None and d["loi"] == []
    assert d["co_the_go"] and d["co_hieu_luc"]
    assert d["bang_chinh"] == "fact_ket_qua_kd"
    assert [b["ma"] for b in d["buoc"]] == ["A1", "A2", "A3", "A4", "A5",
                                            "B1", "B2", "B3", "B4", "B5"]
    assert all(b["trang_thai"] == "ok" for b in d["buoc"])
    buoc = {b["ma"]: b for b in d["buoc"]}
    assert buoc["B1"]["ket_qua"] == ("Đọc lại tệp độc lập rồi so: 7.561 dòng trong tệp = "
                                     "7.561 dòng đã ghi.")
    assert buoc["B2"]["ket_qua"] == "So với số dòng sẽ ghi ở A3: 7.471 = 7.471 dòng."
    assert buoc["B3"]["ket_qua"] == "7.471 = 7.471 dòng."
    assert buoc["B4"]["ket_qua"] == (
        "Lệch 0 ở mọi tháng, mọi cột tiền; số tháng, khách hàng, nhóm, khoản mục khớp. "
        "Tổng cả bảng lệch 0 ở 2/2 bảng (Chi phí 122.689.472.100,65; Kết quả kinh doanh "
        "3.193.114.446.744,50).")
    assert buoc["B5"]["ket_qua"] == "Dữ liệu tháng 1–7 năm 2026 có hiệu lực từ lúc này."
    assert [b["ket_luan"] for b in d["buoc"][5:]] == ["Khớp", "Khớp", "Khớp", "Đúng",
                                                       "Đã chốt"]
    assert admin("GET", "/loads/999999").status_code == 404


def test_doi_chieu_hqkd_5_the(admin, mot_lan):
    the = admin("GET", f"/loads/{mot_lan}/reconcile").json()["the"]
    assert [t["ma"] for t in the] == ["so_dong", "tong", "tien_theo_nhom", "ma", "truoc_sau"]
    assert [t["tieu_de"] for t in the] == [
        "Đối chiếu số dòng: tệp gốc ↔ database",
        "Đối chiếu tổng tiền cả bảng: tệp gốc ↔ database",
        "Đối chiếu tổng tiền theo tháng, cột tiền, nhóm",
        "Đối chiếu tháng, khách hàng, nhóm, khoản mục",
        "Các tháng trong năm 2026 trước và sau lần nạp",
    ]
    so_dong = the[0]
    assert so_dong["tong"] == ["Tổng", 7561, 7561, 7471, 7471, 90, None]
    kh = so_dong["dong"][2]
    assert kh["o"][0] == "Danh mục khách hàng" and kh["phu"] == "dim_khach_hang"
    assert kh["ket_luan_phu"] == "Đủ sau khi bỏ 90 dòng trùng y hệt"
    assert kh["o"][1] == {"v": 227, "mo": {"sheet": 2}}
    assert kh["o"][3]["mo"] == {"bang": "dim_khach_hang", "lop": "silver"}
    assert [d["o"][1:] for d in the[1]["dong"]] == [
        [122689472100.6453, 122689472100.6453, 0],
        [3193114446744.495, 3193114446744.495, 0]]
    tien = the[2]
    assert tien["tuy_chon"][0]["gia_tri"] == "fact_ket_qua_kd"
    assert tien["tuy_chon"][1]["gia_tri"] == "doanh_thu"
    assert tien["dong"][0]["o"][:4] == ["Tháng 1", 50, 50, 196631400595]
    assert tien["dong"][2]["mo"] == {"bang": "fact_ket_qua_kd", "lop": "gold", "nam": 2026,
                                     "ky": 3}
    ma = {d["o"][0]: d for d in the[3]["dong"]}
    assert ma["Số khách hàng trong Kết quả kinh doanh"]["o"][1:] == [69, 69]
    assert ma["Số khoản mục trong Chi phí"]["o"][1:] == [20, 20]
    assert ma["Số mã khách hàng trong Danh mục khách hàng"]["o"][1:] == [
        "135 mã (227 dòng)", "135 mã (137 dòng)"]
    thieu = ma["Mã khách hàng có trong số liệu nhưng không có trong Danh mục khách hàng"]
    assert thieu["o"][1:] == ["BP-TTXK-N04 · TARGET AUSTRALIA PTY. LTD. (tháng 6, 7)", 1]
    assert thieu["ket_luan"] == "Ghi nhận"
    truoc_sau = the[4]
    assert len(truoc_sau["dong"]) == 12
    assert truoc_sau["dong"][0]["o"] == ["Tháng 1/2026", "Có", "0 / 0", "50 / 1.000",
                                         "Ghi thêm"]
    assert truoc_sau["dong"][11]["ket_luan"] == "Giữ nguyên"


def test_doi_chieu_mot_the_theo_lua_chon(admin, mot_lan):
    url = (f"/loads/{mot_lan}/reconcile?the=tien_theo_nhom&bang=fact_chi_phi"
           "&cot=so_tien&chia_theo=ma_khoan_cp")
    the = admin("GET", url).json()["the"]
    assert len(the) == 1
    t = the[0]
    assert t["cot"][0]["t"] == "Khoản mục"
    assert [c["v"] for c in t["tuy_chon"][2]["lua_chon"]] == [
        "ky_thang", "ma_nhom_kd", "ma_khoan_cp"]
    assert any(d.get("phu") for d in t["dong"])
    assert t["tong"][1] == 6960 and t["tong"][3] == t["tong"][4]
    # Tổng các nhóm bằng tổng cả bảng (NT-15).
    tong_bang = admin("GET", f"/loads/{mot_lan}/reconcile?the=tong").json()["the"][0]
    assert t["tong"][3] == tong_bang["dong"][0]["o"][1]
    kq = admin("GET", f"/loads/{mot_lan}/reconcile?the=tien_theo_nhom&bang=fact_ket_qua_kd"
                      "&cot=tat_ca").json()["the"][0]
    tong_kqkd = tong_bang["dong"][1]["o"][1]
    assert kq["tong"][3] == pytest.approx(tong_kqkd)
    assert admin("GET", f"/loads/{mot_lan}/reconcile?the=khong_co").status_code == 404


def test_tai_danh_sach_loi_csv(admin, tmp_path):
    so = openpyxl.load_workbook(can_tep_mau())
    del so["DANH MỤC CHI PHÍ"]
    tep = tmp_path / "loi.xlsx"
    so.save(tep)
    lid = tai_len(admin, tep).json()["load_id"]
    r = admin("GET", f"/loads/{lid}/errors.csv")
    assert r.status_code == 200
    assert r.headers["content-disposition"] == f'attachment; filename="loi-lan-nap-{lid}.csv"'
    dong = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert dong[0] == ["Sheet", "Vị trí", "Vấn đề", "Cách xử lý"]
    assert dong[1][:2] == ["DANH MỤC CHI PHÍ", "Toàn bộ sheet"]
    the = admin("GET", f"/loads/{lid}/reconcile").json()["the"]
    assert the == [{"ma": "tu_choi", "tieu_de": "Đối chiếu tệp gốc ↔ database",
                    "thong_bao": "Không có dữ liệu đối chiếu vì tệp bị từ chối ở bước "
                                 "kiểm tra cấu trúc, chưa có gì được ghi vào database.",
                    "cot": [], "dong": []}]


def test_tep_goc(admin, mot_lan):
    r = admin("GET", f"/loads/{mot_lan}/file")
    assert r.status_code == 200 and r.content == can_tep_mau().read_bytes()
    sheet = admin("GET", f"/loads/{mot_lan}/file/sheets").json()
    assert [(s["so"], s["ten"], s["so_dong"]) for s in sheet] == [
        (1, "TỔNG HỢP", 524), (2, "DANH SACH KHACH HANG THEO HDX", 232),
        (3, "DANH MỤC CHI PHÍ", 27)]
    nd = admin("GET", f"/loads/{mot_lan}/file/sheets/3?trang=1&moi=25").json()
    assert nd["tong"] == 27 and len(nd["dong"]) == 25 and nd["dong"][0]["rn"] == 1
    assert set(nd["dong"][0]) == {"rn", "o", "an"}
    assert admin("GET", f"/loads/{mot_lan}/file/sheets/9").status_code == 404


def test_go_lan_khong_moi_nhat_409(admin, loader, sach):
    """QT-17, B18: chỉ gỡ được lần nạp thành công mới nhất; gỡ xong dữ liệu lần
    trước được dùng lại."""
    lan1 = _nap(admin)
    lan2 = _nap(admin)
    assert dem(sach, "SELECT count(*) FROM gold.fact_ket_qua_kd WHERE is_current") == 348
    assert loader("DELETE", f"/loads/{lan2}").status_code == 403
    r = admin("DELETE", f"/loads/{lan1}")
    assert r.status_code == 409
    assert r.json()["detail"] == ("Có 1 lần nạp sau lần này trên cùng nhóm thông tin. "
                                  "Hãy gỡ các lần nạp sau trước.")
    assert admin("GET", f"/loads/{lan1}").json()["co_the_go"] is False
    assert admin("DELETE", f"/loads/{lan2}").status_code == 204
    assert admin("GET", f"/loads/{lan2}").status_code == 404
    assert dem(sach, "SELECT count(*) FROM gold.fact_ket_qua_kd WHERE is_current") == 348
    assert admin("DELETE", f"/loads/{lan1}").status_code == 204
    assert dem(sach, "SELECT count(*) FROM gold.fact_ket_qua_kd") == 0
    with sach.catalog.transaction() as so:
        hanh_dong = [r[0] for r in so.execute(
            "SELECT action FROM ctl_audit_event WHERE object_id IN (?, ?) ORDER BY event_id",
            (str(lan1), str(lan2)))]
    assert hanh_dong.count("load.success") == 2
    assert hanh_dong.count("load.rollback") == 2


def test_xoa_lich_su_lan_bi_tu_choi(admin, sach, tmp_path):
    so = openpyxl.load_workbook(can_tep_mau())
    del so["DANH MỤC CHI PHÍ"]
    tep = tmp_path / "loi.xlsx"
    so.save(tep)
    lid = tai_len(admin, tep).json()["load_id"]
    assert admin("DELETE", f"/loads/{lid}").status_code == 204
    assert dem(sach, "SELECT count(*) FROM ctl.load") == 0
    assert admin("DELETE", f"/loads/{lid}").status_code == 404
