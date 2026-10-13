"""Bộ đọc tệp HQKD (BAO_CAO_HQKH v5): hợp đồng xuất và ba khoản cấu thành giá vốn.

Workbook dựng trong tmp_path theo hằng số của customer_report_layout, mã và tên đều giả.
Reconcile so bộ đọc openpyxl (CustomerReportReader) với bộ đọc XML (CustomerReportVerifyReader),
nên mỗi kiểm tra chạy trên cả hai.
"""

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from open_webui.data_portal.domain.tables import column_purpose
from open_webui.data_portal.registry.loader import load_definitions
from open_webui.data_portal.sources import customer_report_layout as layout
from open_webui.data_portal.sources.customer_report import CustomerReportReader
from open_webui.data_portal.sources.customer_report_verify import CustomerReportVerifyReader

DEFINITIONS_DIR = Path(layout.__file__).resolve().parents[1] / 'definitions'

# Cột sheet TỔNG HỢP: (nhóm chi phí ở dòng 1, mã ở dòng 2, tiêu đề ở dòng 3).
SUMMARY_COLUMNS = [
    (None, None, 'THÁNG'),
    (None, None, 'NHÓM KINH DOANH'),
    (None, None, 'MÃ KHÁCH HÀNG'),
    (None, None, 'Khách hàng'),
    (None, None, 'Hợp đồng xuất'),
    (None, 'DA01', 'Doanh thu'),
    (None, 'CA00', 'Giá vốn'),
    (None, 'CA04', 'CM Gia Công'),
    (None, 'CA01', 'Nguyên phụ liệu'),
    (None, 'CF01', 'Thêu + In + Wash'),
    (None, 'DA02', 'Thu nhập khác'),
    (None, None, 'LỢI NHUẬN NET'),
    (None, None, 'Tỷ lệ lãi/DT'),
    (None, None, 'Lợi Nhuận Gộp (doanh thu-giá vốn)'),
    (None, None, 'Tỷ lệ (LN Gộp/Doanh Thu)'),
    (None, None, 'Tỷ lệ lãi trên vốn'),
    (None, None, 'Tỷ trọng đóng góp doanh thu'),
    (None, None, 'Tỷ trọng đóng góp lợi nhuận ròng'),
    (None, None, 'Mức chênh lệch giữa lợi nhuận gộp và lợi nhuận ròng'),
    ('Chi phí trực tiếp khách hàng', 'CG02', 'Chi phí vận chuyển'),
    (None, None, 'Trích Chi Phí Trả Lương Khối Văn Phòng'),
    (None, None, 'Trích Chi Phí Chung'),
]

# Tháng, nhóm, mã khách, tên, hợp đồng, doanh thu, giá vốn, CA04, CA01, CF01, rồi các cột còn lại.
SUMMARY_ROWS = [
    ['1', 'KD1', 'KH-A', 'Khach A', 'HD-01', 1000, 600, 100, 450, 50],
    ['2', 'KD1', 'KH-A', 'Khach A', 'HD-02', 2000, 1200, 200, 900, 100],  # hợp đồng chỉ có ở TỔNG HỢP
    ['2', 'KD2', 'KH-B', 'Khach B', None, 500, 300, 0, 300, 0],  # không ghi hợp đồng
    ['2', 'KD2', 'KH-D', 'Khach D tong hop', '  HD-09 ', 700, 400, 0, 400, 0],  # mã chưa có trong danh mục
    [None, None, None, 'TỔNG CỘNG', None, 4200, 2500],
    ['LK', None, 'KH-A', 'Khach A', None, 3000, 1800],
]

# Sheet khách hàng: hai cột HỢP ĐỒNG XUẤT (A và H); bộ đọc phải lấy cột H, sau Trung Tâm Phí.
CUSTOMER_HEADER = {'A': 'HỢP ĐỒNG XUẤT', 'B': 'Mã Trung Tâm Phí', 'C': 'Trung Tâm Phí', 'H': 'HỢP ĐỒNG XUẤT'}
CUSTOMER_ROWS = [
    {'A': 'COT-A-1', 'B': 'KH-A', 'C': 'Khach A chuan', 'H': 'HD-01'},
    {'A': 'COT-A-2', 'B': 'KH-C', 'C': 'Khach C'},
    {'A': 'x', 'B': 'x', 'C': 'x', 'H': 'x'},  # dòng đánh dấu, bỏ qua
]


def _build_workbook(path: Path) -> Path:
    workbook = Workbook()
    summary = workbook.active
    summary.title = layout.SUMMARY_SHEET
    for column, (group, code, header) in enumerate(SUMMARY_COLUMNS, start=1):
        summary.cell(1, column, group)
        summary.cell(2, column, code)
        summary.cell(3, column, header)
    for row_number, values in enumerate(SUMMARY_ROWS, start=4):
        for column, value in enumerate(values, start=1):
            summary.cell(row_number, column, value)

    customers = workbook.create_sheet(layout.CUSTOMER_LIST_SHEET)
    customers['A1'] = 'DANH SÁCH KHÁCH HÀNG THEO HỢP ĐỒNG XUẤT'
    for letter, header in CUSTOMER_HEADER.items():
        customers[f'{letter}3'] = header
    for row_number, values in enumerate(CUSTOMER_ROWS, start=4):
        for letter, value in values.items():
            customers[f'{letter}{row_number}'] = value

    cost_items = workbook.create_sheet(layout.COST_ITEM_SHEET)
    cost_items.append(['MÃ KHOẢN MỤC', 'KHOẢN MỤC'])
    cost_items.append(['CG02', 'Chi phí vận chuyển'])
    cost_items.append(['CA04', 'CM Gia Công'])

    workbook.save(path)
    return path


@pytest.fixture(scope='module')
def registry():
    return load_definitions(DEFINITIONS_DIR)


@pytest.fixture
def workbook_path(tmp_path):
    return _build_workbook(tmp_path / 'hqkd.xlsx')


def _reader_table(path: Path, form, table: str) -> list[dict]:
    reader = CustomerReportReader(path, form)
    assert reader.source_issues() == []
    columns = layout.FLAT_TABLES[table]
    return [row.values for row in reader.rows(table, {column: column for column in columns})]


def _verify_table(path: Path, table: str) -> list[dict]:
    result = CustomerReportVerifyReader(path).read_table(table)
    return [{column: cells.get(index) for index, column in result.headers.items()} for cells in result.rows.values()]


def _customers(rows: list[dict]) -> list[tuple]:
    return [(row['ma_khach'], row['ten_khach'], row['hop_dong_xuat']) for row in rows]


EXPECTED_CUSTOMERS = [
    ('KH-A', 'Khach A chuan', 'HD-01'),
    ('KH-C', 'Khach C', None),
    ('KH-A', 'Khach A chuan', 'HD-02'),
    ('KH-D', 'Khach D tong hop', '  HD-09 '),
]


def test_customer_catalog_takes_contract_after_cost_center(registry, workbook_path):
    rows = _reader_table(workbook_path, registry.form('BAO_CAO_HQKH'), layout.CUSTOMERS_TABLE)
    assert _customers(rows) == EXPECTED_CUSTOMERS


def test_verify_reader_builds_the_same_customer_catalog(workbook_path):
    assert _customers(_verify_table(workbook_path, layout.CUSTOMERS_TABLE)) == EXPECTED_CUSTOMERS


@pytest.mark.parametrize('source', ['reader', 'verify'])
def test_cost_of_goods_parts_come_from_their_codes(registry, workbook_path, source):
    if source == 'reader':
        rows = _reader_table(workbook_path, registry.form('BAO_CAO_HQKH'), layout.RESULTS_TABLE)
    else:
        rows = _verify_table(workbook_path, layout.RESULTS_TABLE)
    parts = [
        (row['ma_khach'], row['gia_von'], row['cm_gia_cong'], row['nguyen_phu_lieu'], row['theu_in_wash'])
        for row in rows
    ]
    assert parts == [
        ('KH-A', '600', '100', '450', '50'),
        ('KH-A', '1200', '200', '900', '100'),
        ('KH-B', '300', '0', '300', '0'),
        ('KH-D', '400', '0', '400', '0'),
    ]


def test_missing_contract_column_in_customer_sheet_is_reported(registry, tmp_path):
    path = _build_workbook(tmp_path / 'hqkd.xlsx')
    workbook = load_workbook(path)
    workbook[layout.CUSTOMER_LIST_SHEET]['H3'] = None
    workbook.save(path)

    reader = CustomerReportReader(path, registry.form('BAO_CAO_HQKH'))
    issues = reader.source_issues()
    assert [(issue.code, issue.params['name']) for issue in issues] == [('MISSING_COLUMN', 'HỢP ĐỒNG XUẤT')]


def test_add_summary_contracts_adds_only_new_pairs():
    customers = [
        {'ma_khach': 'KH-A', 'ten_khach': 'Ten danh muc', 'hop_dong_xuat': 'HD-01'},
        {'ma_khach': 'KH-C', 'ten_khach': 'Khach C', 'hop_dong_xuat': None},
    ]
    summary = [
        {'ma_khach': 'KH-A', 'ten_khach': 'Ten tong hop', 'hop_dong_xuat': ' HD-01 '},  # đã có (khác khoảng trắng)
        {'ma_khach': 'KH-A', 'ten_khach': 'Ten tong hop', 'hop_dong_xuat': 'HD-02'},
        {'ma_khach': 'KH-A', 'ten_khach': 'Ten tong hop', 'hop_dong_xuat': 'HD-02'},  # trùng trong TỔNG HỢP
        {'ma_khach': 'KH-B', 'ten_khach': 'Khach B', 'hop_dong_xuat': '   '},  # hợp đồng trống
        {'ma_khach': 'KH-B', 'ten_khach': 'Khach B', 'hop_dong_xuat': None},
        {'ma_khach': 'KH-D', 'ten_khach': 'Khach D', 'hop_dong_xuat': 'HD-09'},
    ]
    layout.add_summary_contracts(customers, summary)
    assert customers[2:] == [
        {'ma_khach': 'KH-A', 'ten_khach': 'Ten danh muc', 'hop_dong_xuat': 'HD-02'},
        {'ma_khach': 'KH-D', 'ten_khach': 'Khach D', 'hop_dong_xuat': 'HD-09'},
    ]


def test_form_v5_keys_customers_by_code_and_contract(registry):
    form = registry.form('BAO_CAO_HQKH')
    assert form.version == 5
    assert registry.table('dim_khach_hang').business_key == ['ma_khach', 'hop_dong_xuat']
    facts = registry.table('fact_ket_qua_kd').column_names
    assert {'cm_gia_cong', 'nguyen_phu_lieu', 'theu_in_wash'} <= set(facts)


def test_fact_customer_column_still_joins_the_composite_key_dim(registry):
    tables = registry.form('BAO_CAO_HQKH').tables
    fact = registry.table('fact_ket_qua_kd')
    purpose = column_purpose(fact, fact.column('ma_khach'), tables)
    assert 'dim_khach_hang' in purpose

    dim = registry.table('dim_khach_hang')
    assert 'hop_dong_xuat' in column_purpose(dim, dim.column('ma_khach'), tables)
