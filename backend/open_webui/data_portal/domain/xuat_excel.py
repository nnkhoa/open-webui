"""Tải Excel ở Chi tiết bảng (CN-43, mục 19.2): các dòng đang lọc kèm sheet THONG_TIN.

Tên cột là **tên biến** ở cả sheet dữ liệu lẫn THONG_TIN (QT-20). Không theo
trang: lấy mọi dòng khớp bộ lọc.
"""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font

from ..registry.schema import FormTable

COT_THONG_TIN = ("ten_cot", "ten_trong_tep_nbc", "kieu_du_lieu", "bat_buoc", "y_nghia",
                 "dung_de", "vi_du")


def ten_tep(ma_nhom: str, table: FormTable, lop: str, nam: int | None) -> str:
    """`{NHÓM}-{bảng}-{lớp}[-{năm}].xlsx`."""
    duoi = f"-{nam}" if nam is not None and table.cot_nam is not None else ""
    return f"{ma_nhom}-{table.name}-{lop}{duoi}.xlsx"


def ghi(table: FormTable, du_lieu: dict, *, nhom: str, ten_lop: str, bo_loc: str,
        cot: list[dict]) -> bytes:
    """`du_lieu` là kết quả `bang.doc(..., moi=None)`; `cot` là `bang.thong_tin_cot`."""
    so = Workbook()
    sheet = so.active
    sheet.title = table.name[:31]
    sheet.append([c.name for c in du_lieu["cot"]])
    for dong in du_lieu["dong"]:
        sheet.append(dong)
    sheet.freeze_panes = "A2"

    tt = so.create_sheet("THONG_TIN")
    for dong in (
        ("Bảng", table.name),
        ("Tên bảng", table.label),
        ("Nội dung", table.description),
        ("Mỗi dòng là", table.grain),
        ("Dùng để", table.purpose),
        ("Nhóm thông tin", nhom),
        ("Lớp dữ liệu", ten_lop),
        ("Bộ lọc", bo_loc),
        ("Số dòng", du_lieu["tong"]),
        ("Xuất lúc", datetime.now().strftime("%Y/%m/%d - %H:%M")),
    ):
        tt.append(dong)
    tt.append([])
    tt.append(list(COT_THONG_TIN))
    dau = tt.max_row
    for c in cot:
        tt.append([c["ten"], c["ten_nbc"], c["kieu_hien"], "Có" if c["bat_buoc"] else "Không",
                   c["y_nghia"], c["dung_de"], c["vi_du"] or ""])

    dam = Font(bold=True)
    for o in tt["A"][:dau]:
        o.font = dam
    for o in tt[dau]:
        o.font = dam
    for o in sheet[1]:
        o.font = dam
    for chu, rong in zip("ABCDEFG", (22, 34, 16, 10, 60, 60, 24), strict=True):
        tt.column_dimensions[chu].width = rong

    bo_dem = io.BytesIO()
    so.save(bo_dem)
    return bo_dem.getvalue()
