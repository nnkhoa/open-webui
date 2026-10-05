from __future__ import annotations

import datetime as dt
import io
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Font

from ..registry.schema import FormTable
from .tables import TableRows

INFO_SHEET = 'THONG_TIN'
LAYER_LABELS = {'gold': 'Dữ liệu phân tích', 'silver': 'Dữ liệu chuẩn hoá', 'bronze': 'Dữ liệu gốc'}
INFO_COLUMNS = ('ten_cot', 'ten_trong_tep_nbc', 'kieu_du_lieu', 'bat_buoc', 'y_nghia', 'dung_de', 'vi_du')
INFO_COLUMN_WIDTHS = (22, 34, 16, 10, 60, 60, 24)
LABEL_TABLE = 'Bảng'
LABEL_TABLE_NAME = 'Tên bảng'
LABEL_DESCRIPTION = 'Nội dung'
LABEL_GRAIN = 'Mỗi dòng là'
LABEL_PURPOSE = 'Dùng để'
LABEL_DOMAIN = 'Nhóm thông tin'
LABEL_LAYER = 'Lớp dữ liệu'
LABEL_FILTERS = 'Bộ lọc'
LABEL_ROW_COUNT = 'Số dòng'
LABEL_EXPORTED_AT = 'Xuất lúc'
REQUIRED_YES = 'Có'
REQUIRED_NO = 'Không'
NO_FILTER = 'Không lọc'
FILTER_YEAR = 'nam'
FILTER_PERIOD = 'ky'
FILTER_QUERY = 'tim'
EXPORTED_AT_FORMAT = '%Y/%m/%d - %H:%M'
SHEET_NAME_MAX_LENGTH = 31


@dataclass
class ExportInfo:
    domain: str
    layer: str
    filters: str


def file_name(domain_code: str, table: FormTable, layer: str, year: int | None) -> str:
    suffix = f'-{year}' if year is not None and table.year_column is not None else ''
    return f'{domain_code}-{table.name}-{layer}{suffix}.xlsx'


def describe_filters(year: int | None, period: str, query: str) -> str:
    filters = ((FILTER_YEAR, year), (FILTER_PERIOD, period), (FILTER_QUERY, query))
    return ', '.join(f'{name} = {value}' for name, value in filters if value) or NO_FILTER


def build_workbook(table: FormTable, data: TableRows, info: ExportInfo, columns: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = table.name[:SHEET_NAME_MAX_LENGTH]
    sheet.append([column.name for column in data.columns])
    for row in data.rows:
        sheet.append(row)
    sheet.freeze_panes = 'A2'
    bold = Font(bold=True)
    for cell in sheet[1]:
        cell.font = bold

    _write_info_sheet(workbook.create_sheet(INFO_SHEET), table, data, info, columns)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_info_sheet(sheet, table: FormTable, data: TableRows, info: ExportInfo, columns: list[dict]) -> None:
    for row in (
        (LABEL_TABLE, table.name),
        (LABEL_TABLE_NAME, table.label),
        (LABEL_DESCRIPTION, table.description),
        (LABEL_GRAIN, table.grain),
        (LABEL_PURPOSE, table.purpose),
        (LABEL_DOMAIN, info.domain),
        (LABEL_LAYER, LAYER_LABELS[info.layer]),
        (LABEL_FILTERS, info.filters),
        (LABEL_ROW_COUNT, data.total),
        (LABEL_EXPORTED_AT, dt.datetime.now().strftime(EXPORTED_AT_FORMAT)),
    ):
        sheet.append(row)
    sheet.append([])
    sheet.append(list(INFO_COLUMNS))
    header_row = sheet.max_row
    for column in columns:
        sheet.append(
            [
                column['name'],
                column['source_name'],
                column['type_label'],
                REQUIRED_YES if column['required'] else REQUIRED_NO,
                column['meaning'],
                column['purpose'],
                column['example'] or '',
            ]
        )

    bold = Font(bold=True)
    for cell in sheet['A'][:header_row]:
        cell.font = bold
    for cell in sheet[header_row]:
        cell.font = bold
    for letter, width in zip('ABCDEFG', INFO_COLUMN_WIDTHS, strict=True):
        sheet.column_dimensions[letter].width = width
