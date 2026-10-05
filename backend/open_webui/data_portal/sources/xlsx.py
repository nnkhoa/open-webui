"""Mở tệp Excel và đổi giá trị ô thành văn bản — dùng chung cho các bộ đọc.

Giá trị ô chuyển thành văn bản đúng như trong tệp, không ép kiểu, không làm
tròn — việc ép kiểu là của lớp chuẩn hoá.
"""

from __future__ import annotations

import zipfile
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

from ..errors import SourceFileError


def o_trong(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def o_thanh_van_ban(value) -> str | None:
    """Giá trị ô Excel thành văn bản, không sửa nội dung. Ô trống là None."""
    if o_trong(value):
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value)


def ten_cac_sheet(path: Path) -> list[str]:
    """Tên mọi sheet của tệp, theo thứ tự trong tệp, kể cả sheet ẩn."""
    so = mo_so(path, read_only=True)
    try:
        return list(so.sheetnames)
    finally:
        so.close()


def mo_so(path: Path, *, data_only: bool = True, read_only: bool = False):
    """Mở sổ Excel, lỗi thì báo câu người dùng hiểu được."""
    try:
        return load_workbook(path, data_only=data_only, read_only=read_only)
    except zipfile.BadZipFile:
        raise SourceFileError(
            "Tệp không phải định dạng .xlsx hợp lệ. Nếu tệp là .xls cũ, hãy mở bằng "
            "Excel và lưu lại thành .xlsx."
        ) from None
    except Exception as exc:
        raise SourceFileError(f"Không mở được tệp: {exc}") from None

